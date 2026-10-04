"""Коннекторы (текст, docx, ASR, Ollama LLM/Vision, выгрузка, папка-приёмник) и темы."""
import io
import json
import os
import time
from datetime import date
from pathlib import Path

import httpx
import pytest

from meetflow.connectors.asr_whisper import WhisperLocalASR
from meetflow.connectors.base import ConnectorError
from meetflow.connectors.export import DocxOutput, MarkdownOutput
from meetflow.connectors.files_docx import DocxSource, docx_to_text
from meetflow.connectors.files_text import TextSource
from meetflow.connectors.llm_ollama import OllamaLLM
from meetflow.connectors.registry import PluginRegistry, default_registry
from meetflow.connectors.vision_ollama import OllamaVision
from meetflow.connectors.watch_folder import MARKER, FolderIntake, check_folder_path
from meetflow.core.config import ThemeError, list_themes, load_theme, validate
from meetflow.core.types import SourceKind

LINES = "\n".join(f"[00:0{i}:00] Игнатов Иван: реплика номер {i} про склад" for i in range(6))


# --- текст -------------------------------------------------------------------
@pytest.mark.parametrize("enc", ["utf-8", "utf-8-sig", "utf-16", "cp1251"])
def test_text_encodings(enc):
    doc = TextSource().read(LINES.encode(enc), "t.txt")
    assert "склад" in doc.text and doc.kind == SourceKind.TRANSCRIPT


@pytest.mark.parametrize("data,msg", [(b"", "пустой"), (b"\x00\x01\x02" * 100, "не похож"), (b"ok", "значащих")])
def test_text_rejects(data, msg):
    with pytest.raises(ConnectorError, match=msg):
        TextSource().read(data, "x.txt")


# --- docx --------------------------------------------------------------------
def _docx_bytes():
    from docx import Document
    d = Document()
    d.add_paragraph("Протокол встречи. Повестка: склад")
    t = d.add_table(rows=1, cols=2)
    cell = t.cell(0, 0)
    cell.text = "освещению"
    cell.add_paragraph("Дата: 10.10")          # второй абзац в ячейке — не склеивать
    run = t.cell(0, 1).paragraphs[0].add_run("строка1")
    run.add_break()
    run.add_text("строка2")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_docx_cells_not_glued_and_breaks_kept():
    text = docx_to_text(_docx_bytes())
    assert "освещениюДата" not in text and "освещению / Дата: 10.10" in text
    assert "строка1\nстрока2" in text


def test_docx_corrupt():
    with pytest.raises(ConnectorError, match="повреждён"):
        DocxSource().read(b"PK not a zip really" * 3, "a.docx")


# --- ASR ---------------------------------------------------------------------
def test_asr_offline_env_and_waiting_tools(monkeypatch, tmp_path):
    assert os.environ.get("HF_HUB_OFFLINE") == "1"
    asr = WhisperLocalASR()
    monkeypatch.setattr(asr, "available", lambda: False)
    f = tmp_path / "rec 01.10.2026.mp4"
    f.write_bytes(b"x")
    with pytest.raises(ConnectorError, match="waiting_tools"):
        asr.transcribe(f)
    assert f.exists()  # запись не теряется


def test_asr_leaves_media_for_pipeline_to_remove(monkeypatch, tmp_path):
    """Перевёрнут в итерации 4: коннектор больше не удаляет запись сам —
    удаляет конвейер после сохранения расшифровки (находка ревью №22)."""
    class Seg:
        def __init__(self, s, t): self.start, self.text = s, t

    class FakeModel:
        def transcribe(self, *a, **k):
            return [Seg(1, "Игнатов, подготовьте отчёт"), Seg(65, " ")], type("I", (), {"duration": 1800})()

    asr = WhisperLocalASR(duration_min=20, duration_max=90)
    asr._model = FakeModel()
    f = tmp_path / "rec.mp4"
    f.write_bytes(b"x")
    doc = asr.transcribe(f)
    assert doc.kind == SourceKind.TRANSCRIPT and "[00:00:01]" in doc.text and f.exists()
    assert doc.meta["media_path"] == str(f)


def test_pipeline_removes_media_after_saving(store, project, meeting, make_pipe, tmp_path):
    from meetflow.connectors.asr_whisper import WhisperLocalASR

    class Seg:
        start, text = 1, "Петрова Анастасия, подготовьте регламент склада"

    p = make_pipe()
    asr = p.plugins.get("asr", "whisper_local")
    asr._model = type("M", (), {"transcribe": lambda *a, **k: ([Seg()], type("I", (), {"duration": 1800})())})()
    f = tmp_path / "rec 01.10.2026.mp4"
    f.write_bytes(b"x")
    r = p.ingest(project, meeting, f.name, path=f)
    assert r["status"] == "ok" and not f.exists()
    assert store.get("sources", r["source_id"])["removed_reason"] == "запись распознана, файл удалён"


def test_asr_duration_window(tmp_path):
    asr = WhisperLocalASR(duration_min=20, duration_max=90)
    asr._model = type("M", (), {"transcribe": lambda *a, **k: ([], type("I", (), {"duration": 600})())})()
    f = tmp_path / "rec.mp4"
    f.write_bytes(b"x")
    with pytest.raises(ConnectorError, match="held_recording"):
        asr.transcribe(f)


# --- Ollama ------------------------------------------------------------------
def _transport(responses):
    calls = []

    def handler(req: httpx.Request):
        calls.append(json.loads(req.content or b"{}"))
        code, body = responses.pop(0) if responses else (200, {"message": {"content": "{}"}})
        return httpx.Response(code, json=body)
    return httpx.MockTransport(handler), calls


def test_ollama_json_mode_and_ctx():
    tr, calls = _transport([(200, {"message": {"content": '{"tasks": []}'}, "prompt_eval_count": 10, "eval_count": 3})])
    r = OllamaLLM(model="glm", transport=tr, num_ctx=65536, backoff=0).complete("s", "u")
    assert r.text == '{"tasks": []}' and calls[0]["format"] == "json" and calls[0]["options"]["num_ctx"] == 65536


def test_ollama_retries_5xx_then_fails():
    tr, calls = _transport([(500, {}), (503, {}), (500, {})])
    with pytest.raises(ConnectorError, match="после 3 попыток"):
        OllamaLLM(model="glm", transport=tr, backoff=0).complete("s", "u")
    assert len(calls) == 3


def test_ollama_refuses_remote_url():
    with pytest.raises(ConnectorError, match="вне локального контура"):
        OllamaLLM(url="http://10.0.0.5:11434")


def test_ollama_health_reports_missing_model():
    tr, _ = _transport([(200, {"models": [{"name": "other"}]})])
    h = OllamaLLM(model="glm", transport=tr).health()
    assert h["ok"] is False and "не загружена" in h["detail"]


def test_vision_describe_is_material_not_fact():
    tr, calls = _transport([(200, {"message": {"content": "На доске: сроки релиза"}})])
    doc = OllamaVision(model="dsv", transport=tr).describe(b"\x89PNG....", "доска.png")
    assert doc.kind == SourceKind.OTHER and "images" in calls[0]["messages"][0]


# --- выгрузка ----------------------------------------------------------------
def test_docx_export_no_markdown_artifacts():
    from docx import Document
    data = DocxOutput().render("# Заголовок\n\n**жирный** и `код`\n\n| a | b |\n|---|---|\n| 1 | 2 |\n", "t")
    d = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "**" not in text and "`" not in text and len(d.tables) == 1 and d.tables[0].cell(1, 1).text == "2"
    assert d.styles["Normal"].font.name == "Times New Roman"


def test_md_export_roundtrip():
    assert MarkdownOutput().render("# x", "t") == "# x".encode()


# --- папка-приёмник ------------------------------------------------------------
def test_watch_folder_path_safety(tmp_path):
    with pytest.raises(ConnectorError):
        check_folder_path(tmp_path, home=tmp_path)
    busy = tmp_path / "busy"
    busy.mkdir()
    (busy / "x.txt").write_text("1")
    with pytest.raises(ConnectorError, match="метки"):
        check_folder_path(busy, home=tmp_path / "home")
    (busy / MARKER).write_text("")
    check_folder_path(busy, home=tmp_path / "home")


def test_watch_folder_settle_date_dedup(tmp_path):
    (tmp_path / MARKER).write_text("")
    a = tmp_path / "Расшифровка 01.10.2026.txt"
    a.write_text(LINES, encoding="utf-8")
    nd = tmp_path / "без даты.txt"
    nd.write_text(LINES + " ", encoding="utf-8")
    fi = FolderIntake(tmp_path, today=date(2026, 10, 3))
    r1 = fi.scan()
    assert not r1.ready and len(r1.waiting) == 2          # первый проход — ещё «оседает»
    r2 = fi.scan()
    assert [i.meeting_date for i in r2.ready] == [date(2026, 10, 1)] and r2.no_date == ["без даты.txt"]
    assert fi.scan().no_date == []                         # о файле без даты — один раз
    fi.known.add(r2.ready[0].digest)
    copy = tmp_path / "копия 01.10.2026.txt"
    copy.write_text(LINES, encoding="utf-8")
    fi.scan()
    r4 = fi.scan()
    assert "копия 01.10.2026.txt" in r4.duplicates and (tmp_path / "обработано" / "копия 01.10.2026.txt").exists()


def test_watch_folder_unavailable(tmp_path):
    with pytest.raises(ConnectorError, match="недоступна"):
        FolderIntake(tmp_path / "нет").scan()


# --- реестр плагинов ------------------------------------------------------------
def test_registry_resolves_by_extension():
    r = default_registry()
    assert r.resolve_source("a.docx").name == "docx"
    assert r.resolve_source("a.MP3").name == "whisper_local"
    assert r.resolve_source("a.jpg").name == "ollama_vision"
    assert r.resolve_source("a.exe") is None


def test_registry_rejects_duplicate_name():
    r = PluginRegistry()
    r.register(TextSource())
    with pytest.raises(ConnectorError):
        r.register(TextSource())


# --- темы --------------------------------------------------------------------
@pytest.mark.parametrize("name", list_themes())
def test_all_themes_load(name):
    th = load_theme(name)
    assert th.word("protocol") and th["protocol_skeleton"]


def test_themes_differ_in_vocabulary_and_palette():
    a, b, c = load_theme("erp_rollout"), load_theme("daily_ops"), load_theme("it_product")
    assert {a.word("meeting"), b.word("meeting"), c.word("meeting")} == {"встреча", "планёрка", "синк"}
    assert {a["design"]["palette"], b["design"]["palette"], c["design"]["palette"]} == {"indigo", "forest", "polaris"}
    assert "director" in b["roles"] and "director" not in a["roles"]


def test_theme_todo_lists_placeholders():
    assert any("llm.generator.model" in t for t in load_theme("erp_rollout").todo())


@pytest.mark.parametrize("patch,err", [
    ({"design": {"palette": "neon"}}, "palette"),
    ({"roles": ["member"]}, "super_admin"),
    ({"protocol_skeleton": []}, "protocol_skeleton"),
    ({"llm": {"generator": {"model": "m"}, "verifier": {"model": "m"}}}, "совпадает"),
    ({"llm": {"generator": {"model": "x:contributor"}, "verifier": {"model": "y"}}}, "обучением"),
])
def test_theme_validation(patch, err):
    base = dict(load_theme("_default").data)
    base.update(patch)
    assert any(err in e for e in validate(base))


def test_theme_missing_file():
    with pytest.raises(ThemeError):
        load_theme("нет_такой")


def test_theme_from_custom_dir(tmp_path):
    src = Path(__file__).resolve().parents[1] / "themes" / "_default.json"
    (tmp_path / "_default.json").write_text(src.read_text(encoding="utf-8"), encoding="utf-8")
    (tmp_path / "school.json").write_text(json.dumps({"id": "school", "title": "Школа",
        "vocabulary": {"meeting": "педсовет"}}, ensure_ascii=False), encoding="utf-8")
    th = load_theme("school", tmp_path)
    assert th.word("meeting") == "педсовет" and th.word("protocol") == "Протокол встречи"
