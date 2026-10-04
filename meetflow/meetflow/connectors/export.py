"""Выгрузка документов: markdown и Word.

Протокол хранится как markdown (источник истины), .docx — производное.
Уроки md_to_docx: не оставлять `**` и обратные кавычки (≈60 файлов были
испорчены), markdown-таблицы — настоящие таблицы Word, заголовки — стилями.
Шрифт документа берётся из темы (design.document_font), по умолчанию
Times New Roman 12 pt — корпоративный стиль владельца.
"""
from __future__ import annotations

import io
import re

from .base import OutputConnector

_INLINE = re.compile(r"(\*\*|__|`)")


class MarkdownOutput(OutputConnector):
    name = "md"
    extension = ".md"

    def render(self, markdown: str, title: str) -> bytes:
        return markdown.encode("utf-8")


class DocxOutput(OutputConnector):
    name = "docx"
    extension = ".docx"

    def __init__(self, font: str = "Times New Roman", size_pt: int = 12):
        self.font, self.size = font, size_pt

    def render(self, markdown: str, title: str) -> bytes:
        from docx import Document
        from docx.shared import Pt

        doc = Document()
        st = doc.styles["Normal"]
        st.font.name, st.font.size = self.font, Pt(self.size)
        table_rows: list[list[str]] = []

        def flush_table():
            if not table_rows:
                return
            rows = [r for r in table_rows if not all(set(c) <= set("-: ") for c in r)]
            t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            t.style = "Table Grid"
            for i, r in enumerate(rows):
                for j, c in enumerate(r):
                    t.cell(i, j).text = _INLINE.sub("", c)
            table_rows.clear()

        for line in markdown.splitlines():
            s = line.strip()
            if s.startswith("|") and s.endswith("|"):
                table_rows.append([c.strip() for c in s.strip("|").split("|")])
                continue
            flush_table()
            if not s:
                continue
            m = re.match(r"^(#{1,4})\s+(.*)$", s)
            if m:
                doc.add_heading(_INLINE.sub("", m.group(2)), level=min(len(m.group(1)), 4))
            elif s.startswith(("- ", "* ")):
                doc.add_paragraph(_INLINE.sub("", s[2:]), style="List Bullet")
            else:
                doc.add_paragraph(_INLINE.sub("", s))
        flush_table()
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
